"""Excel-Dateien sicher öffnen und als Werteraster lesen.

Excel-Dateien sind nicht vertrauenswürdig: keine Makros, keine Formelauswertung, Größenprüfung
vor dem Entpacken. openpyxl nutzt automatisch defusedxml, wenn installiert (geprüft beim Import).
"""

from __future__ import annotations

import datetime as dt
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl
import openpyxl.xml

from backend.excel.numbers import PercentCell

if not openpyxl.xml.DEFUSEDXML:  # pragma: no cover - Absicherung der Installation
    raise RuntimeError("defusedxml fehlt: Excel-Dateien würden ungeschützt geparst")

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"


class ExcelRejected(Exception):
    """Datei wird abgelehnt. Die Meldung ist für Benutzer gedacht."""


@dataclass
class SheetData:
    name: str
    rows: list[list]  # Werte, 0-basiert; Zeile i entspricht Excel-Zeile i+1
    merged: list[tuple[int, int, int, int]]  # (min_row, min_col, max_row, max_col), 1-basiert
    formula_without_cache: set[tuple[int, int]] = field(default_factory=set)  # (row, col) 1-basiert
    hidden: bool = False


def check_file(path: Path, original_name: str, max_uncompressed: int, max_entries: int) -> None:
    name = original_name.lower()
    with open(path, "rb") as fh:
        head = fh.read(8)
    if name.endswith(".xls") or head.startswith(OLE_MAGIC):
        raise ExcelRejected(
            "Altes Excel-Format (.xls) wird nicht unterstützt. Bitte in Excel als .xlsx speichern."
        )
    if name.endswith((".xlsm", ".xltm", ".xlsb")):
        raise ExcelRejected(
            "Dateien mit Makros (.xlsm/.xlsb) werden nicht angenommen. "
            "Bitte in Excel als .xlsx (ohne Makros) speichern."
        )
    if not name.endswith(".xlsx"):
        raise ExcelRejected("Nur .xlsx-Dateien werden unterstützt.")
    if not head.startswith(ZIP_MAGIC):
        raise ExcelRejected("Die Datei ist keine gültige .xlsx-Datei.")
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
            if len(infos) > max_entries:
                raise ExcelRejected("Die Datei enthält ungewöhnlich viele Bestandteile.")
            total = sum(i.file_size for i in infos)
            if total > max_uncompressed:
                raise ExcelRejected(
                    f"Die Datei ist entpackt zu groß ({total // 1_000_000} MB, "
                    f"erlaubt {max_uncompressed // 1_000_000} MB)."
                )
            names = {i.filename.lower() for i in infos}
            if any(n.endswith("vbaproject.bin") for n in names):
                raise ExcelRejected("Die Datei enthält Makros und wird nicht angenommen.")
            if "[content_types].xml" not in names:
                raise ExcelRejected("Die Datei ist keine gültige .xlsx-Datei.")
            ctypes = zf.read("[Content_Types].xml")
            if b"macroEnabled" in ctypes:
                raise ExcelRejected("Die Datei ist als Makro-Datei gekennzeichnet.")
    except zipfile.BadZipFile as exc:
        raise ExcelRejected("Die Datei ist beschädigt oder keine gültige .xlsx-Datei.") from exc


def _convert(cell) -> object:
    value = cell.value
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        fmt = cell.number_format or ""
        if "%" in fmt:
            base = Decimal(value) if isinstance(value, int) else Decimal(repr(value))
            return PercentCell(base * 100)
        return value
    if isinstance(value, (dt.datetime, dt.date)):
        return value
    if isinstance(value, dt.time):
        return value.isoformat()
    return str(value)


def list_sheets(path: Path) -> list[str]:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ExcelRejected("Die Datei konnte nicht gelesen werden (beschädigt?).") from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_sheet(path: Path, sheet: str | None = None, max_rows: int | None = None) -> SheetData:
    """Liest ein Blatt vollständig (Werte, verbundene Zellen, Formeln ohne Ergebnis)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, keep_links=False, keep_vba=False)
    except Exception as exc:
        raise ExcelRejected("Die Datei konnte nicht gelesen werden (beschädigt?).") from exc
    try:
        name = sheet or wb.sheetnames[0]
        if name not in wb.sheetnames:
            raise ExcelRejected(f"Tabellenblatt {name!r} nicht gefunden.")
        ws = wb[name]
        rows: list[list] = []
        for row in ws.iter_rows():
            if max_rows is not None and len(rows) >= max_rows:
                break
            rows.append([_convert(c) for c in row])
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        merged = [
            (r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges
        ]
        hidden = ws.sheet_state != "visible"
    finally:
        wb.close()

    formula_cells = _formula_cells(path, name, max_rows)
    missing = {
        (r, c)
        for (r, c) in formula_cells
        if r - 1 < len(rows) and (c - 1 >= len(rows[r - 1]) or rows[r - 1][c - 1] is None)
    }
    return SheetData(name=name, rows=rows, merged=merged, formula_without_cache=missing, hidden=hidden)


def _formula_cells(path: Path, sheet: str, max_rows: int | None) -> set[tuple[int, int]]:
    """Zellen mit Formel (zweiter Durchlauf ohne data_only, nur Streaming)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    result = set()
    try:
        ws = wb[sheet]
        for r_idx, row in enumerate(ws.iter_rows(), start=1):
            if max_rows is not None and r_idx > max_rows:
                break
            for c_idx, cell in enumerate(row, start=1):
                if getattr(cell, "data_type", None) == "f":
                    result.add((r_idx, c_idx))
    finally:
        wb.close()
    return result
