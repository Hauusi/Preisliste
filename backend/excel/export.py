"""Excel-Export eines Vergleichs bzw. einer Kalkulation.

Schutz gegen Formel-Injektion: Jeder Text wird ausdrücklich als Text geschrieben (nie als Formel).
Texte, die mit = + - @ oder Tab/CR beginnen, werden zusätzlich markiert, damit sie auch beim
späteren Bearbeiten in Excel nicht als Formel gelten. Beträge werden als exakte Dezimalzahlen geschrieben.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from backend.comparison.service import STATUS_LABELS, STATUSES, percent_change
from backend.services.manufacturers import code_map, with_code
from backend.models.entities import (
    CalculationResult,
    CalculationRun,
    Comparison,
    ComparisonItem,
    ImportMessage,
    Manufacturer,
    PriceList,
    Rule,
)

DANGEROUS_PREFIX = ("=", "+", "-", "@", "\t", "\r")
MONEY_FORMAT = "#,##0.00"
PERCENT_FORMAT = '0.00" %"'
BOLD = Font(bold=True)


class _Sheet:
    def __init__(self, wb, title: str, header: list[str], widths: list[int] | None = None):
        self.ws = wb.create_sheet(title)
        for i, w in enumerate(widths or [], start=1):
            self.ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        self.ws.freeze_panes = "A2"
        self.row([self._cell(h, bold=True) for h in header])

    def _cell(self, value, fmt: str | None = None, bold: bool = False):
        c = WriteOnlyCell(self.ws, value=value)
        if isinstance(value, str):
            c.data_type = "s"  # nie als Formel
            if value.startswith(DANGEROUS_PREFIX):
                c.alignment = Alignment(horizontal="left")
                c.quotePrefix = True  # Excel behandelt den Inhalt dauerhaft als Text
        if fmt:
            c.number_format = fmt
        if bold:
            c.font = BOLD
        return c

    def row(self, values: list) -> None:
        self.ws.append([v if isinstance(v, Cell) else self._cell(v) for v in values])

    def money(self, v):
        return self._cell(v, MONEY_FORMAT) if isinstance(v, Decimal) else self._cell(v)

    def percent(self, v):
        return self._cell(v, PERCENT_FORMAT) if isinstance(v, Decimal) else self._cell(v)


ITEM_HEADER = ["Artikelnummer", "Original-Nr.", "Alte Artikelnummer", "Hersteller", "Bezeichnung", "Kategorie", "Staffel ab",
               "Preis alt", "Preis neu", "Währung", "Differenz", "Differenz %", "Status", "Zuordnung", "Hinweis"]
ITEM_WIDTHS = [18, 16, 18, 18, 40, 18, 10, 12, 12, 9, 12, 12, 16, 14, 50]


def _item_row(sh: _Sheet, i: ComparisonItem, mfr: dict, codes: dict) -> None:
    code = codes.get(i.manufacturer_id)
    old_nr = with_code(i.old_article.article_number, codes.get(i.old_article.manufacturer_id)) if i.old_article else None
    new_nr = with_code(i.article_number, code)
    hint = " ".join(filter(None, [
        i.note,
        f"Geändert: {', '.join(i.changed_fields)}" if i.changed_fields else None,
        ("Kandidaten: " + ", ".join(f"{with_code(c['number'], code)} ({c['score']})" for c in i.candidates))
        if i.candidates else None,
    ]))
    sh.row([new_nr, i.article_number, old_nr if old_nr != new_nr else None, mfr.get(i.manufacturer_id),
            i.description, i.category, i.min_quantity, sh.money(i.old_amount), sh.money(i.new_amount), i.currency,
            sh.money(i.difference), sh.percent(i.difference_percent), STATUS_LABELS.get(i.status, i.status),
            i.match_method, hint or None])


def _calc_sheet(wb, db: Session, run: CalculationRun | None, title: str = "Kalkulation") -> None:
    sh = _Sheet(wb, title, ["Zeile", "Artikelnummer", "Original-Nr.", "Bezeichnung", "Ausgangspreis", "Ergebnis",
                            "Währung", "Status", "Fehler", "Rechenweg"], [8, 18, 16, 40, 14, 14, 9, 10, 40, 90])
    codes = code_map(db)
    if run is None:
        sh.row(["Keine Kalkulation für die neue Liste vorhanden."])
        return
    rule = db.get(Rule, run.rule_version.rule_id)
    sh.row([f"Regel {rule.name} v{run.rule_version.version}, Menge {run.quantity}, berechnet {run.created_at:%d.%m.%Y %H:%M} UTC"])
    stmt = (select(CalculationResult).where(CalculationResult.run_id == run.id)
            .options(selectinload(CalculationResult.article)).order_by(CalculationResult.id))
    for r in db.scalars(stmt):
        steps = []
        for t in r.trace:
            part = f"{t.get('nr')}. {t.get('text') or t.get('typ')}"
            if t.get("operand"):
                part += f" {t['operand']}"
            if t.get("basis_wert"):
                part += f" (Basis {t['basis_wert']})"
            if t.get("betrag"):
                part += f" Betrag {t['betrag']}"
            part += f" = {t.get('ergebnis')}"
            steps.append(part)
        sh.row([r.article.source_row, with_code(r.article.article_number, codes.get(r.article.manufacturer_id)),
                r.article.article_number, r.article.description, sh.money(r.start_amount),
                sh.money(r.result_amount), r.currency, r.status, r.error, " | ".join(steps)])


def export_comparison(db: Session, cmp: Comparison) -> bytes:
    mfr = {m.id: m.name for m in db.scalars(select(Manufacturer))}
    codes = code_map(db)
    items = db.scalars(select(ComparisonItem).where(ComparisonItem.comparison_id == cmp.id)
                       .options(selectinload(ComparisonItem.old_article)).order_by(ComparisonItem.id)).all()
    wb = openpyxl.Workbook(write_only=True)

    s = cmp.summary or {}
    summary = _Sheet(wb, "Zusammenfassung", ["Angabe", "Wert"], [34, 60])
    for label, value in [
        ("Alte Liste", f"{cmp.old_list.name} ({cmp.old_list.source_file})"),
        ("Neue Liste", f"{cmp.new_list.name} ({cmp.new_list.source_file})"),
        ("Hersteller", mfr.get(cmp.manufacturer_id, "alle")),
        ("Verglichener Preis", cmp.price_type + (f" (Regel-Version v{cmp.rule_version.version})" if cmp.rule_version else "")),
        ("Menge", cmp.quantity),
        ("Berechnet (UTC)", f"{(cmp.updated_at or cmp.created_at):%d.%m.%Y %H:%M}"),
        ("Analysiert", s.get("analysiert", 0)),
    ] + [(STATUS_LABELS[st], s.get(st, 0)) for st in STATUSES]:
        summary.row([label, value])

    def items_sheet(title: str, wanted) -> None:
        sh = _Sheet(wb, title, ITEM_HEADER, ITEM_WIDTHS)
        for i in items:
            if wanted(i):
                _item_row(sh, i, mfr, codes)

    items_sheet("Alle Artikel", lambda i: True)
    items_sheet("Preisänderungen", lambda i: i.status in ("PREIS_ERHOEHT", "PREIS_GESENKT"))
    items_sheet("Neue Artikel", lambda i: i.status == "NEUER_ARTIKEL")
    items_sheet("Entfallene Artikel", lambda i: i.status == "ENTFALLENER_ARTIKEL")
    items_sheet("Unklare Zuordnungen", lambda i: i.status == "NICHT_EINDEUTIG")

    run = db.scalar(select(CalculationRun).where(CalculationRun.price_list_id == cmp.new_price_list_id)
                    .options(selectinload(CalculationRun.rule_version)).order_by(CalculationRun.id.desc()).limit(1))
    _calc_sheet(wb, db, run)

    errors = _Sheet(wb, "Fehler", ["Quelle", "Zeile", "Artikelnummer", "Spalte", "Code", "Meldung"],
                    [22, 8, 18, 18, 22, 80])
    for i in items:
        if i.status == "FEHLER":
            errors.row(["Vergleich", None, with_code(i.article_number, codes.get(i.manufacturer_id)), None,
                        "VERGLEICH", i.note])
    for label, list_id in (("Import alte Liste", cmp.old_price_list_id), ("Import neue Liste", cmp.new_price_list_id)):
        for m in db.scalars(select(ImportMessage).where(ImportMessage.price_list_id == list_id,
                                                        ImportMessage.level == "FEHLER")
                            .order_by(ImportMessage.source_row)):
            errors.row([label, m.source_row, None, m.column, m.code, m.text])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_calculation(db: Session, run: CalculationRun) -> bytes:
    wb = openpyxl.Workbook(write_only=True)
    pl = db.get(PriceList, run.price_list_id)
    info = _Sheet(wb, "Zusammenfassung", ["Angabe", "Wert"], [30, 60])
    rule = db.get(Rule, run.rule_version.rule_id)
    for label, value in [("Preisliste", f"{pl.name} ({pl.source_file})"), ("Regel", f"{rule.name} v{run.rule_version.version}"),
                         ("Menge", run.quantity), ("Artikel", (run.summary or {}).get("artikel")),
                         ("Berechnet", (run.summary or {}).get("OK")), ("Fehler", (run.summary or {}).get("FEHLER"))]:
        info.row([label, value])
    _calc_sheet(wb, db, run)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_price_update(db: Session, upd, draft: bool = False) -> bytes:
    """Jahresabgleich: neue EK/VK-Liste mit unseren Artikeln. Entwurf deutlich gekennzeichnet."""
    from backend.comparison.update import REASON_LABELS, STATUS_LABELS as UL
    from backend.models.entities import PriceUpdateItem, User

    codes = code_map(db)
    users = {u.id: u.username for u in db.scalars(select(User))}
    items = db.scalars(select(PriceUpdateItem).where(PriceUpdateItem.update_id == upd.id)
                       .order_by(PriceUpdateItem.id)).all()
    rule_label = None
    if upd.rule_version:
        rule_label = f"{db.get(Rule, upd.rule_version.rule_id).name} v{upd.rule_version.version}"
    wb = openpyxl.Workbook(write_only=True)
    s = upd.summary or {}
    open_ = s.get("offen", 0)
    info = _Sheet(wb, "Zusammenfassung", ["Angabe", "Wert"], [36, 70])
    if draft:
        info.row([info._cell("ENTWURF", bold=True),
                  info._cell(f"ENTWURF – {open_} Positionen ungeprüft. Nicht als Preisliste verwenden.", bold=True)])
    basis = (f"UVP/RRP, EK = UVP - {upd.dealer_discount} % Händlerrabatt" if upd.price_type == "UVP"
             else "EK direkt")
    if upd.list_currency and upd.list_currency != "EUR":
        basis += f", Währung {upd.list_currency} × Kurs {upd.exchange_rate} = EUR (auf Cent gerundet)"
    from backend.comparison.update import EXCEPTION_LABELS
    exceptions = "; ".join(f"{EXCEPTION_LABELS.get(ex['typ'], ex['typ'])} {ex['wert']}: {ex['regel']}"
                           for ex in upd.exceptions or []) or "keine"
    for label, value in [
        ("Unsere Liste (EK/VK bisher)", f"{upd.base_list.name} ({upd.base_list.source_file})"),
        ("Herstellerliste (neu)", f"{upd.source_list.name} ({upd.source_list.source_file})"),
        ("Umfang", "Teilliste (nur einzelne Serien)" if upd.scope == "TEIL" else "Jahrespreisliste"),
        ("Herstellerliste liefert", basis), ("VK-Standardregel (aus neuem EK)", rule_label or "keine"),
        ("Eigene Kalkulation (Ausnahmen)", exceptions),
        ("Gegenrechnung", "jeder VK unabhängig nachgerechnet; Abweichungen sind als Fehler markiert"),
        ("Prüfschwelle EK-Änderung %", upd.review_threshold), ("Artikel", s.get("artikel", 0)),
    ] + [(UL[k], s.get(k, 0)) for k in UL] + [
        ("Noch ungeprüft", open_),
        ("Nur beim Hersteller (ignoriert)", s.get("ignoriert_nur_beim_hersteller", 0))]:
        info.row([label, value])

    decisions = {"NEU": "neuer Preis", "ALT": "alter Preis behalten", "MANUELL": "VK von Hand"}
    checks = {True: "stimmt", False: "ABWEICHUNG"}
    cols = [("Artikelnummer", 18), ("Original-Nr.", 16), ("Bezeichnung", 40), ("Herstellerpreis", 14),
            ("Währung Hersteller", 10), ("EK alt", 12), ("EK neu", 12), ("EK Δ%", 10), ("VK alt", 12),
            ("VK neu", 12), ("VK Δ%", 10), ("Faktor VK/EK", 11), ("Marge %", 10), ("Währung", 9), ("Status", 20),
            ("Prüfhinweise", 60), ("Entscheidung", 20), ("Regel", 36), ("Rechenweg EK", 50),
            ("Gegenrechnung", 13), ("geprüft von", 14), ("geprüft am (UTC)", 18)]
    header, widths = [c[0] for c in cols], [c[1] for c in cols]
    main = _Sheet(wb, "ENTWURF Neue Preisliste" if draft else "Neue Preisliste", header, widths)
    open_items = _Sheet(wb, "Zu prüfen", header, widths)
    for i in items:
        hints = "; ".join([REASON_LABELS.get(r, r) for r in i.reasons or []] + ([i.note] if i.note else []))
        new = i.decision != "ALT"
        ek_pct = percent_change(i.old_amount, i.final_ek) if new and i.old_amount and i.final_ek is not None else None
        vk_pct = percent_change(i.vk_old, i.final_vk) if new and i.vk_old and i.final_vk is not None else None
        row = [with_code(i.article_number, codes.get(i.manufacturer_id)), i.article_number, i.description,
               main.money(i.source_amount), i.source_currency, main.money(i.old_amount), main.money(i.final_ek),
               main.percent(ek_pct), main.money(i.vk_old), main.money(i.final_vk), main.percent(vk_pct),
               main._cell(i.factor, "0.0000") if i.factor is not None and i.decision == "NEU" else None,
               main.percent(margin(i.final_ek, i.final_vk)),
               i.currency, UL.get(i.status, i.status), hints, decisions.get(i.decision, i.decision),
               i.rule_label, i.ek_text, checks.get(i.check_ok, "nicht prüfbar" if i.decision == "NEU" else ""), users.get(i.reviewed_by),
               i.reviewed_at.strftime("%d.%m.%Y %H:%M") if i.reviewed_at else None]
        main.row(row)
        if i.needs_review and not i.reviewed_at:
            open_items.row([open_items._cell(v.value, v.number_format) if isinstance(v, Cell) else v for v in row])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def margin(ek, vk):
    """Rohertragsmarge in % vom VK: (VK − EK) / VK × 100."""
    if ek is None or vk in (None, Decimal(0)):
        return None
    return ((vk - ek) / vk * 100).quantize(Decimal("0.1"))


def export_customer_changes(db: Session, upd) -> bytes:
    """Preisänderungsliste für Kunden: nur Artikel mit geändertem VK, alt -> neu, gültig ab."""
    from backend.models.entities import PriceUpdateItem

    codes = code_map(db)
    items = db.scalars(select(PriceUpdateItem).where(PriceUpdateItem.update_id == upd.id)
                       .order_by(PriceUpdateItem.article_number)).all()
    who = upd.base_list.manufacturer.name if upd.base_list.manufacturer else upd.base_list.name
    valid = ".".join(reversed(upd.valid_from.split("-"))) if upd.valid_from else None
    wb = openpyxl.Workbook(write_only=True)
    sh = _Sheet(wb, "Preisänderungen", ["Artikelnummer", "Bezeichnung", "Preis bisher", "Preis neu", "Änderung %",
                                         "Gültig ab"], [18, 44, 14, 14, 12, 12])
    for i in items:
        if i.vk_old is None or i.final_vk is None or i.final_vk == i.vk_old:
            continue
        change = ((i.final_vk - i.vk_old) / i.vk_old * 100).quantize(Decimal("0.01")) if i.vk_old else None
        sh.row([with_code(i.article_number, codes.get(i.manufacturer_id)), i.description, sh.money(i.vk_old),
                sh.money(i.final_vk), sh.percent(change), valid])
    info = _Sheet(wb, "Info", ["Angabe", "Wert"], [30, 60])
    info.row(["Hersteller", who])
    info.row(["Gültig ab", valid or "nicht angegeben"])
    info.row(["Preise", "netto, in " + (upd.base_list.currency or "EUR")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
