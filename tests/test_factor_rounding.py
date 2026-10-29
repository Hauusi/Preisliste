"""VK-Rundung neuer Hersteller (10 Cent aufrunden) und Faktor VK/EK mit einer Nachkommastelle."""

from decimal import Decimal as D

from tests.test_e2e_szenarien import items_of, new_manufacturer, smart_page, start


def test_new_manufacturer_rounds_vk_up_to_10_cent_also_for_exceptions(admin_client, app):
    c = admin_client
    new_manufacturer(c, "Raphi LED", "RA", "2,8")
    _, page = smart_page(c, "RaphiLED_Preisliste_2026.xlsx")
    start(c, app, page)
    _, page = smart_page(c, "RaphiLED_Herstellerliste_2027.xlsx")
    assert "VK neu = EK × 2,8, auf 10 Cent (aufrunden)" in page
    target = start(c, app, page)
    items = items_of(int(target.rsplit("/", 1)[1]))
    bl2 = items["BL2"]  # EK 3,12 × 2,8 = 8,736 -> 8,74 -> 8,80
    assert (bl2.final_ek, bl2.final_vk, bl2.factor) == (D("3.12"), D("8.80"), D("2.8"))
    for i in items.values():
        if i.decision == "NEU":
            assert i.final_vk % D("0.1") == 0 and i.factor == i.factor.quantize(D("0.1")), i.article_number
            assert i.check_ok is True, i.article_number
    # „anders kalkulieren“ mit Faktor übernimmt die VK-Rundung des Herstellers
    c.post(f"{target}/ausnahmen", data={"csrf_token": c.csrf, "action": "setzen", "factor": "3",
                                        "item": str(bl2.id)})
    bl2 = items_of(int(target.rsplit("/", 1)[1]))["BL2"]
    assert (bl2.final_vk, bl2.factor, bl2.check_ok) == (D("9.40"), D("3.0"), True)  # 9,36 -> 9,40


def test_implausible_ek_keeps_old_price_and_change_is_shown_as_zero(admin_client, tmp_path, app):
    """EK-Sprung (z. B. EAN statt Preis gelesen) wird nie übernommen; unveränderte Preise zeigen 0 % statt leer."""
    import io

    import openpyxl

    from tests.test_price_update import _import, _items, _ours, _setup

    c = admin_client
    mid, rule_id = _setup(c)
    ours = _ours(c, app, tmp_path, mid, [["A", "a", "100,00", "200,00"], ["B", "b", "50,00", "100,00"],
                                         ["C", "c", "10,00", "20,00"]])
    new = _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["A", "a", "45766674"],
                                          ["B", "b", "52,00"]], mid, "HERSTELLER")
    c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new, "rule_id": rule_id})
    upd, items = _items()
    a = items["A"]
    assert (a.status, a.decision, a.final_ek, a.final_vk) == ("FEHLER", "ALT", D("100.00"), D("200.00"))
    assert "EK_UNPLAUSIBEL" in a.reasons and "Zeile 2" in a.note and "45.766.674,00" in a.note
    assert "Herstellerliste Zeile 3" in items["B"].ek_text
    # Sammelbestätigung nimmt den unplausiblen Preis nie mit
    c.post(f"/aktualisierungen/{upd.id}/alle-bestaetigen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    _, items = _items()
    assert items["A"].reviewed_at is None and items["A"].final_ek == D("100.00")
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/aktualisierungen/{upd.id}/export?entwurf=1").content))
    ws = wb["ENTWURF Neue Preisliste"]
    head = [x.value for x in ws[1]]
    rows = {r[1]: dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert (rows["A"]["EK Δ%"], rows["A"]["VK Δ%"]) == (0, 0)  # alter Preis bleibt: 0 %, nicht leer
    assert (rows["C"]["EK Δ%"], rows["C"]["VK Δ%"]) == (0, 0)  # fehlt beim Hersteller
