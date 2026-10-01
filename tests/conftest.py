import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.api.app import create_app
from backend.auth.core import create_user, ip_limiter
from backend.config import Settings
from backend.database import engine as db_engine

ADMIN_PW = "admin-passwort-123"
USER_PW = "benutzer-passwort-123"


@pytest.fixture
def settings(tmp_path):
    return Settings(data_dir=tmp_path / "data", cookie_secure=True)


@pytest.fixture
def app(settings):
    ip_limiter.reset()
    application = create_app(settings)
    with db_engine.session_scope() as db:
        create_user(db, "admin", ADMIN_PW, "admin", settings)
        create_user(db, "anna", USER_PW, "benutzer", settings)
    return application


@pytest.fixture
def client(app):
    return TestClient(app, base_url="https://testserver")


def login(client, username="admin", password=ADMIN_PW):
    r = client.post("/login", data={"username": username, "password": password}, follow_redirects=False)
    assert r.status_code == 303, r.text
    return csrf_of(client)


def csrf_of(client) -> str:
    page = client.get("/konto").text
    return re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)


@pytest.fixture
def admin_client(client):
    client.csrf = login(client)
    return client


@pytest.fixture
def user_client(app):
    c = TestClient(app, base_url="https://testserver")
    c.csrf = login(c, "anna", USER_PW)
    return c


def make_xlsx(path: Path, rows, sheet="Preise", merged=(), number_formats=None, extra_sheets=None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(list(r))
    for rng in merged:
        ws.merge_cells(rng)
    for coord, fmt in (number_formats or {}).items():
        ws[coord].number_format = fmt
    for name, srows in (extra_sheets or {}).items():
        s = wb.create_sheet(name)
        for r in srows:
            s.append(list(r))
    wb.save(path)
    return path


def strip_formula_cache(path: Path) -> None:
    """Formel-Ergebnisse aus der Datei entfernen (wie bei Dateien, die nie in Excel gespeichert wurden)."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = re.sub(rb"(<f>[^<]*</f>)<v>[^<]*</v>", rb"\1", data)
            dst.writestr(item, data)
    tmp.replace(path)
