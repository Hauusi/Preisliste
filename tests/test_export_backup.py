import io
import tarfile
import zipfile

import openpyxl
from sqlalchemy import select

from backend.database.engine import session_scope
from backend.models.entities import Comparison
from backend.services.backup import create_backup
from tests.test_group_b_web import RULE_FORM, import_list


def _flow(admin_client, tmp_path):
    old = import_list(admin_client, tmp_path, "alt", [
        ["Art.-Nr.", "Bezeichnung", "Listenpreis"],
        ["A-1", "Schraube", "100,00"], ["GONE", "Weg", "1,00"], ["ABC-1000", "Spezialteil Typ X", "5,00"],
        ["ZERO", "Null", "0"],
    ])
    new = import_list(admin_client, tmp_path, "neu", [
        ["Art.-Nr.", "Bezeichnung", "Listenpreis"],
        ["A-1", "Schraube", "110,00"], ["NEW", "Neu", "2,00"],
        ["ABC-1001", "Spezialteil Typ X", "6,00"], ["ZERO", "Null", "1"], ["BAD", "-x", "kaputt"],
    ])
    # Gefährliche Texte, wie sie als Textzellen in einer Lieferantenliste stehen könnten
    from backend.models.entities import Article
    with session_scope() as db:
        for a in db.scalars(select(Article).where(Article.price_list_id == new)):
            if a.article_number == "A-1":
                a.description = '=HYPERLINK("http://boese")'
            if a.article_number == "NEW":
                a.description = "+cmd|' /C calc'!A0"
    r = admin_client.post("/regeln/neu", data={**RULE_FORM, "csrf_token": admin_client.csrf}, follow_redirects=False)
    rule_id = int(r.headers["location"].rsplit("/", 1)[1])
    admin_client.post(f"/listen/{new}/kalkulation", data={"csrf_token": admin_client.csrf, "rule_id": rule_id, "quantity": "1"})
    r = admin_client.post("/vergleiche", data={"csrf_token": admin_client.csrf, "old_id": old, "new_id": new,
                                               "price_type": "LISTE", "quantity": "1"}, follow_redirects=False)
    return int(r.headers["location"].rsplit("/", 1)[1])


def test_comparison_export_sheets_values_and_injection(admin_client, tmp_path):
    cmp_id = _flow(admin_client, tmp_path)
    r = admin_client.get(f"/vergleiche/{cmp_id}/export")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Zusammenfassung", "Alle Artikel", "Preisänderungen", "Neue Artikel",
                             "Entfallene Artikel", "Unklare Zuordnungen", "Kalkulation", "Fehler"]
    rows = {row[0]: row for row in wb["Alle Artikel"].iter_rows(min_row=2, values_only=True)}
    assert rows["A-1"][8] == 110 and rows["A-1"][10] == 10 and rows["A-1"][11] == 10
    new_rows = {r[0]: r for r in wb["Neue Artikel"].iter_rows(min_row=2, values_only=True)}
    assert set(new_rows) == {"NEW", "BAD"}
    assert new_rows["BAD"][8] is None and "Kein gültiger LISTE-Preis" in new_rows["BAD"][14]
    assert [r[0] for r in wb["Entfallene Artikel"].iter_rows(min_row=2, values_only=True)] == ["GONE"]
    assert [r[0] for r in wb["Unklare Zuordnungen"].iter_rows(min_row=2, values_only=True)] == ["ABC-1001"]
    kalk = list(wb["Kalkulation"].iter_rows(min_row=3, values_only=True))
    assert any(r[1] == "A-1" and r[5] == 97.24 and "Transport" in r[9] for r in kalk)
    fehler = list(wb["Fehler"].iter_rows(min_row=2, values_only=True))
    assert any(r[0] == "Vergleich" and r[2] == "ZERO" for r in fehler)
    assert any(r[0] == "Import neue Liste" and r[4] == "KEIN_BETRAG" for r in fehler)

    # Formel-Injektion: keine einzige Formel in der Datei, gefährliche Texte als Text mit quotePrefix
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/"):
                assert "<f>" not in z.read(name).decode()
        assert 'quotePrefix="1"' in z.read("xl/styles.xml").decode()
    assert rows["A-1"][4] == '=HYPERLINK("http://boese")'
    assert rows["NEW"][4] == "+cmd|' /C calc'!A0"


def test_calculation_export(admin_client, tmp_path):
    _flow(admin_client, tmp_path)
    page = admin_client.get("/listen/2/kalkulation").text
    run_url = page.split('href="/kalkulationen/')[1].split('"')[0]
    r = admin_client.get(f"/kalkulationen/{run_url}/export")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Zusammenfassung", "Kalkulation"]


def test_export_requires_login(client):
    assert client.get("/vergleiche/1/export", follow_redirects=False).status_code == 303


def test_backup_consistent_and_rotates(admin_client, tmp_path, settings):
    _flow(admin_client, tmp_path)
    paths = [create_backup(settings, keep=2)[0] for _ in range(3)]
    remaining = sorted((settings.data_dir / "backups").glob("preisliste-*.tar.gz"))
    assert len(remaining) <= 2 and paths[-1] in remaining
    with tarfile.open(paths[-1]) as tar:
        names = tar.getnames()
        assert "preisliste.sqlite3" in names and any(n.startswith("uploads/") for n in names)
        tar.extract("preisliste.sqlite3", tmp_path / "restore", filter="data")
    import sqlite3
    con = sqlite3.connect(tmp_path / "restore" / "preisliste.sqlite3")
    assert con.execute("select count(*) from articles").fetchone()[0] == 9
    con.close()


def test_reset_data(admin_client, tmp_path, settings, monkeypatch):
    _flow(admin_client, tmp_path)
    from backend import cli
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    assert cli.main(["reset-data"]) == 2  # ohne --ja passiert nichts
    assert cli.main(["reset-data", "--ja"]) == 0
    import sqlite3
    con = sqlite3.connect(settings.db_path)
    assert con.execute("select count(*) from users").fetchone()[0] == 0
    assert con.execute("select count(*) from articles").fetchone()[0] == 0
    con.close()
    assert list(settings.upload_dir.iterdir()) == []
    assert list((settings.data_dir / "backups").glob("preisliste-*.tar.gz"))  # Sicherung vorhanden


def test_clear_data_keeps_users_and_rules(admin_client, tmp_path, settings, monkeypatch):
    _flow(admin_client, tmp_path)
    from backend import cli
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    assert cli.main(["clear-data"]) == 2
    assert cli.main(["clear-data", "--ja"]) == 0
    import sqlite3
    con = sqlite3.connect(settings.db_path)
    q = lambda t: con.execute(f"select count(*) from {t}").fetchone()[0]
    assert q("manufacturers") == q("price_lists") == q("articles") == q("comparisons") == 0
    assert q("users") == 2 and q("rules") == 1 and q("rule_versions") == 1
    con.close()
    assert admin_client.get("/").status_code == 200  # Anmeldung bleibt gültig
