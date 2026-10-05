"""Jahresabgleich: Fremdwährung, Serien-Ausnahmen, Teillisten, Gegenrechnung, Übernahme als aktuelle Liste."""

import io
import re
from decimal import Decimal

import openpyxl
from sqlalchemy import select

from backend.database.engine import session_scope
from backend.models.entities import Article, ArticlePrice, PriceList, RuleException
from tests.test_price_update import _import, _items, _ours, _setup

D = Decimal


def _rule(c, name, factor, mid=None, steps=None):
    data = {"csrf_token": c.csrf, "name": name, "start_price": "EK", "rounding_mode": "HALF_UP",
            "rounding_places": "2", "rounding_timing": "STEP", "step_1_type": "multiply", "step_1_value": factor}
    if mid:
        data["manufacturer_id"] = str(mid)
    data.update(steps or {})
    r = c.post("/regeln/neu", data=data, follow_redirects=False)
    assert r.status_code == 303, r.text
    return r.headers["location"].rsplit("/", 1)[1]


def _start(c, ours, new, rule_id, scope="VOLL"):
    r = c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new,
                                          "rule_id": rule_id, "scope": scope}, follow_redirects=False)
    assert r.status_code == 303, r.text
    return r.headers["location"]


SEK = {"list_currency": "SEK", "exchange_rate": "0,095"}


def test_sek_list_is_converted_before_calculation(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c, factor="2,6", **SEK)
    ours = _ours(c, app, tmp_path, mid, [["LED1", "Modul 1", "11,50", "29,90"]])
    new = _import(c, app, tmp_path, "sek", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "Modul 1", "123,45"]],
                  mid, "HERSTELLER", currency="SEK")
    page = c.get(f"/listen/{new}/kalkulation").text
    assert "Währung <b>SEK</b> → EUR mit Kurs <b>0,095</b>" in page
    url = _start(c, ours, new, rule_id)
    upd, items = _items()
    i = items["LED1"]
    # 123,45 SEK × 0,095 = 11,72775 -> 11,73 EUR; × 2,6 = 30,498 -> 30,50
    assert (i.source_amount, i.source_currency, i.final_ek, i.final_vk) == (D("123.45"), "SEK", D("11.73"), D("30.50"))
    assert i.ek_text == "EK Hersteller 123,45 SEK × Kurs 0,095 = 11,73 EUR"
    assert i.check_ok is True and i.status == "OK" and i.currency == "EUR"
    assert upd.list_currency == "SEK" and upd.exchange_rate == D("0.095")
    assert "Gegenrechnung stimmt" in c.get(url).text
    wb = openpyxl.load_workbook(io.BytesIO(c.get(url + "/export").content))
    ws = wb["Neue Preisliste"]
    row = dict(zip([h.value for h in ws[1]], next(ws.iter_rows(min_row=2, values_only=True))))
    assert (row["Herstellerpreis"], row["Währung Hersteller"], row["EK neu"], row["VK neu"]) == (123.45, "SEK", 11.73, 30.5)
    assert row["Gegenrechnung"] == "stimmt" and row["Faktor VK/EK"] == 2.6002


def test_currency_mismatch_is_blocked(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c, **SEK)
    ours = _ours(c, app, tmp_path, mid, [["LED1", "Modul 1", "10,00", "20,00"]])
    # Herstellerliste versehentlich als EUR importiert: wäre Faktor ~10 zu teuer
    new = _import(c, app, tmp_path, "eur", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "Modul 1", "105,00"]],
                  mid, "HERSTELLER")
    assert "Diese Liste ist in EUR importiert" in c.get(f"/listen/{new}/kalkulation").text
    r = c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new, "rule_id": rule_id})
    assert r.status_code == 400 and "in EUR importiert" in r.text and "liefert der Hersteller in SEK" in r.text


def test_manufacturer_currency_needs_rate(admin_client):
    c = admin_client
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "name": "Svensk", "list_currency": "SEK"})
    assert r.status_code == 400 and "Umrechnungskurs fehlt" in r.text


def test_exceptions_on_manufacturer_page(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c, factor="2,6")
    r27 = _rule(c, "RaphiLED 2,7", "2,7", mid)
    r28 = _rule(c, "RaphiLED LED5", "2,8", mid)
    ours = _ours(c, app, tmp_path, mid, [["LED1", "Standard", "10,00", "26,00"], ["NV1", "Nova", "10,00", "27,00"],
                                         ["LED51", "a", "10,00", "28,00"], ["LED52", "b", "10,00", "27,00"],
                                         ["X9", "c", "10,00", "26,00"]])
    new = _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "a", "10,00"],
                                          ["NV1", "b", "10,00"], ["LED51", "c", "10,00"], ["LED52", "d", "10,00"],
                                          ["X9", "e", "10,00"]], mid, "HERSTELLER")
    edit = f"/hersteller/{mid}"
    add = lambda t, v, rid: c.post(edit + "/ausnahmen", data={"csrf_token": c.csrf, "match_type": t, "value": v,
                                                              "rule_id": rid})
    assert "auf 1 Artikelnummer(n), z. B. RANV1" in add("ARTIKEL", "nv1", r27).text
    r = add("PREFIX", "RALED5", r28)
    assert "ACHTUNG: passt bisher auf keinen Artikel" in r.text and "ohne Kürzel RA" in r.text
    with session_scope() as db:
        bad = db.scalar(select(RuleException).where(RuleException.value == "RALED5"))
    r = c.post(f"{edit}/ausnahmen/{bad.id}/loeschen", data={"csrf_token": c.csrf})
    assert r.status_code == 400 and "sicher" in r.text
    c.post(f"{edit}/ausnahmen/{bad.id}/loeschen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    assert "auf 2 Artikelnummer(n)" in add("PREFIX", "led5", r28).text
    r = add("PREFIX", "LED-5", r28)
    assert r.status_code == 400 and "gibt es schon" in r.text
    assert add("SERIE", "Nova", r27).status_code == 400  # Serien werden nicht benannt
    add("ARTIKEL", "LED52", r27)   # einzelner Artikel geht vor Nummernanfang
    add("PREFIX", "X", r27)
    add("PREFIX", "X9", r28)       # X9 passt auf zwei Bereiche mit verschiedenen Regeln

    url = _start(c, ours, new, rule_id)
    upd, items = _items()
    assert items["LED1"].final_vk == D("26.00") and items["LED1"].rule_label.startswith("Standard: RaphiLED v1")
    assert items["NV1"].final_vk == D("27.00") and items["NV1"].rule_label == "Ausnahme Artikel „nv1“: RaphiLED 2,7 v1"
    assert items["LED51"].final_vk == D("28.00") and "Nummer beginnt mit „led5“" in items["LED51"].rule_label
    assert items["LED52"].final_vk == D("27.00")
    x9 = items["X9"]
    assert x9.status == "FEHLER" and "Mehrere Ausnahmen passen" in x9.note
    assert (x9.final_ek, x9.final_vk, x9.decision) == (D("10.00"), D("26.00"), "ALT")
    assert len(upd.exceptions) == 5
    assert "eigene Kalkulation" in c.get(url).text

    # Regel löschen entfernt ihre Ausnahmen (mit Hinweis auf der Bestätigungsseite)
    assert "Ausnahme(n) (Artikel mit eigener Kalkulation) mit dieser Regel werden entfernt" in c.get(f"/regeln/{r28}/loeschen").text
    c.post(f"/regeln/{r28}/loeschen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    with session_scope() as db:
        assert db.scalar(select(RuleException).where(RuleException.rule_id == int(r28))) is None


def test_partial_list_keeps_others_unchanged(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c)
    ours = _ours(c, app, tmp_path, mid, [["LED1", "a", "10,00", "20,00"], ["LED2", "b", "10,00", "20,00"]])
    new = _import(c, app, tmp_path, "teil", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "a", "10,50"]],
                  mid, "HERSTELLER")
    url = _start(c, ours, new, rule_id, scope="TEIL")
    upd, items = _items()
    assert items["LED2"].status == "UNVERAENDERT" and not items["LED2"].needs_review
    assert (items["LED2"].final_ek, items["LED2"].final_vk) == (D("10.00"), D("20.00"))
    assert items["LED1"].final_vk == D("21.00") and upd.summary["offen"] == 0 and upd.scope == "TEIL"
    assert c.get(url + "/export").status_code == 200


def test_adopt_as_current_list(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c)
    ours = _ours(c, app, tmp_path, mid, [["LED1", "a", "10,00", "20,00"], ["LED2", "b", "10,00", "20,00"]])
    new = _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "a", "10,50"]],
                  mid, "HERSTELLER")
    url = _start(c, ours, new, rule_id)
    r = c.post(url + "/uebernehmen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    assert "Noch 1 Positionen ungeprüft" in r.text
    c.post(url + "/alle-bestaetigen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    r = c.post(url + "/uebernehmen", data={"csrf_token": c.csrf})
    assert "Häkchen" in r.text
    r = c.post(url + "/uebernehmen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    assert "Als aktuelle Liste" in r.text
    upd, _ = _items()
    with session_scope() as db:
        pl = db.get(PriceList, upd.adopted_list_id)
        assert pl.kind == "UNSERE" and pl.status == "IMPORTIERT" and "Stand" in pl.name
        prices = {(a.article_number, p.price_type): p.amount for a in db.scalars(
            select(Article).where(Article.price_list_id == pl.id)) for p in a.prices}
    assert prices == {("LED1", "EK"): D("10.50"), ("LED1", "LISTE"): D("21.00"),
                      ("LED2", "EK"): D("10.00"), ("LED2", "LISTE"): D("20.00")}
    r = c.post(url + "/uebernehmen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    assert "bereits als aktuelle Liste übernommen" in r.text
    # nächste Herstellerliste: die übernommene Liste ist vorausgewählt
    new2 = _import(c, app, tmp_path, "h2", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "a", "11,00"]],
                   mid, "HERSTELLER")
    assert re.search(rf'<option value="{pl.id}" selected>', c.get(f"/listen/{new2}/kalkulation").text)


def test_check_mismatch_is_error_and_not_bulk_accepted(admin_client, tmp_path, app, monkeypatch):
    import backend.comparison.update as upd_mod
    from backend.calculations.engine import calculate as real

    def broken(rule, prices, quantity):
        res = real(rule, prices, quantity)
        res.result += D("0.01")  # simulierter Engine-Fehler
        return res

    monkeypatch.setattr(upd_mod, "calculate", broken)
    c = admin_client
    mid, rule_id = _setup(c)
    ours = _ours(c, app, tmp_path, mid, [["LED1", "a", "10,00", "20,00"]])
    new = _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "a", "10,00"]],
                  mid, "HERSTELLER")
    url = _start(c, ours, new, rule_id)
    _, items = _items()
    i = items["LED1"]
    assert i.status == "FEHLER" and i.check_ok is False and "Gegenrechnung abweichend" in i.note
    assert (i.final_ek, i.final_vk) == (D("10.00"), D("20.00"))
    r = c.post(url + "/alle-bestaetigen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    assert "0 Positionen bestätigt" in r.text


def test_formula_rule_is_marked_not_checkable(admin_client, tmp_path, app):
    c = admin_client
    mid, _ = _setup(c)
    rid = _rule(c, "Formel", "1", mid, {"step_2_type": "formula", "step_2_value": "current * 2"})
    ours = _ours(c, app, tmp_path, mid, [["LED1", "a", "10,00", "20,00"]])
    new = _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "a", "10,00"]],
                  mid, "HERSTELLER")
    _start(c, ours, new, rid)
    _, items = _items()
    i = items["LED1"]
    assert i.final_vk == D("20.00") and i.status == "PRUEFEN" and i.reasons == ["NICHT_GEGENGEPRUEFT"]
    assert i.check_ok is None
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/aktualisierungen/{i.update_id}/export?entwurf=1").content))
    ws = wb["ENTWURF Neue Preisliste"]
    row = dict(zip([h.value for h in ws[1]], next(ws.iter_rows(min_row=2, values_only=True))))
    assert row["Gegenrechnung"] == "nicht prüfbar"


def test_manufacturer_pages_render(admin_client, user_client):
    c = admin_client
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "name": "Redtronic", "code": "RT"}, follow_redirects=False)
    edit = r.headers["location"].split("?")[0]
    page = c.get(r.headers["location"]).text
    assert "Angelegt" in page and "Artikel mit eigener Kalkulation" in page and 'name="exchange_rate"' in page
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "id": edit.rsplit("/", 1)[1], "name": "Redtronic",
                                    "code": "RT", "list_basis": "EK", "list_currency": "SEK",
                                    "exchange_rate": "0,095", "review_threshold": "10"})
    assert "Gespeichert" in r.text
    assert "in SEK (× 0,095)" in c.get("/hersteller").text
    # fremder Hersteller ist für andere Benutzer nicht sichtbar
    assert user_client.get(edit).status_code == 404
    assert 'href="/hersteller/' not in user_client.get("/hersteller").text
