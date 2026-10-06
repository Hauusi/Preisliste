"""End-to-End: kompletter Arbeitsablauf mit den Testdateien aus beispiel-daten/ (wie ein Mitarbeiter klickt).

Jeder Preis wird unabhängig nachgerechnet: EK aus der Datei (ggf. UVP − Rabatt, × Kurs), VK = EK × Faktor.
"""

import io
import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
from sqlalchemy import select

from backend.database.engine import session_scope
from backend.models.entities import Article, Job, Manufacturer, PriceList, PriceUpdate, PriceUpdateItem
from tests.conftest import run_jobs
from tests.test_import_web import upload

D = Decimal
DATA = Path(__file__).resolve().parent.parent / "beispiel-daten"


def money(v) -> Decimal:
    return D(v).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def new_manufacturer(c, name, code, factor, **settings) -> int:
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "name": name, "code": code, "factor": factor},
               follow_redirects=False)
    assert r.status_code == 303, r.text
    mid = int(r.headers["location"].split("?")[0].rsplit("/", 1)[1])
    if settings:
        base = {"csrf_token": c.csrf, "id": str(mid), "name": name, "code": code, "factor": factor,
                "list_basis": "EK", "review_threshold": "10", "ignore_leading_zeros": "1"}
        r = c.post("/hersteller", data={**base, **settings})
        assert "Gespeichert" in r.text, r.text
    return mid


def smart_page(c, file: str):
    r = upload(c, DATA / file)
    assert r.status_code == 303, r.text
    url = r.headers["location"]
    return url, c.get(url).text


def start_form(page: str) -> dict:
    form = re.search(r'<form method="post" action="(/import/\d+/start)">(.*?)</form>', page, re.S)
    fields = dict(re.findall(r'<input type="hidden" name="(\w+)" value="([^"]*)">', form.group(2)))
    return form.group(1), fields


def start(c, app, page: str, **extra) -> str:
    """Startet den Import wie der Button auf der Prüfseite und folgt der automatischen Weiterleitung."""
    action, fields = start_form(page)
    r = c.post(action, data={**fields, **extra}, follow_redirects=False)
    assert r.status_code == 303, r.text
    job_url = r.headers["location"]
    run_jobs(app)
    job_page = c.get(job_url).text
    target = re.search(r'http-equiv="refresh" content="0;url=([^"]+)"', job_page)
    assert target, job_page
    return target.group(1)


def items_of(upd_id: int) -> dict:
    with session_scope() as db:
        return {i.article_number: i for i in db.scalars(select(PriceUpdateItem)
                                                        .where(PriceUpdateItem.update_id == upd_id))}


def file_prices(file: str, header_row: int, col: int, strip: str = "") -> dict:
    ws = openpyxl.load_workbook(DATA / file, data_only=True).active
    out = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[0]:
            out[str(row[0])[len(strip):] if strip and str(row[0]).startswith(strip) else str(row[0])] = row[col]
    return out


def test_raphi_led_full_year(admin_client, app):
    c = admin_client
    mid = new_manufacturer(c, "Raphi LED", "RA", "2,6")

    # 1) Unsere Liste 2026: alles automatisch erkannt
    url, page = smart_page(c, "RaphiLED_Preisliste_2026.xlsx")
    assert f'<option value="{mid}" selected>Raphi LED (RA)</option>' in page and "automatisch erkannt" in page
    assert '<option value="UNSERE" selected>' in page and 'name="strip_code" value="1" checked' in page
    assert "<b>RALED1</b>" in page and "(RALED1)" not in page  # Anzeige mit Kürzel, gespeichert ohne
    assert "Als unsere Liste importieren" in page and "disabled" not in page.split("Als unsere Liste importieren")[0][-200:]
    target = start(c, app, page)
    assert target.startswith("/listen/") and target.endswith("?neu=1")
    ours_id = int(target.split("/")[2].split("?")[0])
    assert "aktuelle EK/VK-Liste von Raphi LED" in c.get(target).text
    with session_scope() as db:
        arts = db.scalars(select(Article).where(Article.price_list_id == ours_id)).all()
        assert len(arts) == 50 and all(a.status == "OK" for a in arts) and arts[0].article_number == "LED1"

    # 2) Neue Herstellerliste 2027: Hersteller, Art, Abdeckung erkannt; Abgleich startet automatisch
    url, page = smart_page(c, "RaphiLED_Herstellerliste_2027.xlsx")
    assert f'<option value="{mid}" selected>' in page and '<option value="HERSTELLER" selected>' in page
    assert "<b>47 von 50</b>" in page and "3 fehlen in der Liste" in page and "4 neue des Herstellers" in page
    assert "VK neu = EK × 2,6" in page and '<option value="VOLL" selected>' in page
    target = start(c, app, page)
    assert re.fullmatch(r"/aktualisierungen/\d+", target)
    upd_id = int(target.rsplit("/", 1)[1])
    items = items_of(upd_id)
    assert len(items) == 50
    new = file_prices("RaphiLED_Herstellerliste_2027.xlsx", 4, 2)
    checked = 0
    for nr, i in items.items():
        if nr in new and i.status in ("OK", "PRUEFEN"):
            ek = money(str(new[nr]))
            assert i.final_ek == ek and i.final_vk == money(ek * D("2.6")), nr
            assert i.check_ok is True
            checked += 1
    assert checked == 47
    big = {k for k, i in items.items() if i.status == "PRUEFEN"}
    assert big == {"LEUCHTBALKEN3", "AS300", "BL3"}
    assert all(items[k].reasons == ["EK_AENDERUNG"] for k in big)
    assert {k for k, i in items.items() if i.status == "NICHT_IN_HERSTELLERLISTE"} == {"RK20", "ZB3"}
    renamed = items["AS500"]
    assert renamed.status == "NICHT_EINDEUTIG" and renamed.candidates[0]["number"] == "AS500-B"
    page = c.get(target).text
    assert "6 von 6 Positionen noch zu prüfen" in page and "RALEUCHTBALKEN3" in page

    # 3) Prüfen: geänderte Nummer zuordnen, Rest per Sammelbestätigung, dann ein Klick „Abschließen“
    c.post(f"{target}/positionen/{renamed.id}", data={"csrf_token": c.csrf, "action": "zuordnen",
                                                      "new_id": str(renamed.candidates[0]["new_id"])})
    r = c.post(f"{target}/alle-bestaetigen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    assert "Positionen bestätigt" in r.text
    items = items_of(upd_id)
    assert items["AS500"].final_ek == money(str(new["AS500-B"])) and items["AS500"].match_method == "BESTAETIGT"
    page = c.get(target).text
    assert "bereit zum Abschließen" in page
    r = c.post(f"{target}/abschliessen", data={"csrf_token": c.csrf}, follow_redirects=False)
    assert r.headers["location"] == f"{target}?fertig=1"
    page = c.get(r.headers["location"]).text
    assert "Fertig!" in page and f'content="1;url={target}/export"' in page
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"{target}/export").content))
    rows = {r[0]: r for r in wb["Neue Preisliste"].iter_rows(min_row=2, values_only=True)}
    assert len(rows) == 50 and "RALED1" in rows and "RARK20" in rows
    with session_scope() as db:
        upd = db.get(PriceUpdate, upd_id)
        current = db.get(PriceList, upd.adopted_list_id)
        assert current.kind == "UNSERE"
        led1 = db.scalar(select(Article).where(Article.price_list_id == current.id, Article.article_number == "LED1"))
        prices = {p.price_type: p.amount for p in led1.prices}
    assert prices == {"EK": money(str(new["LED1"])), "LISTE": money(money(str(new["LED1"])) * D("2.6"))}

    # 4) Preiserhöhung unterm Jahr: nur Arbeitsscheinwerfer -> Teilliste vorgeschlagen, Rest unverändert
    url, page = smart_page(c, "RaphiLED_Preiserhoehung_Arbeitsscheinwerfer.xlsx")
    assert "<b>4 von 50</b>" in page and '<option value="TEIL" selected>' in page
    assert f"in „{current.name}“" in page  # Grundlage ist die gerade übernommene Liste
    target = start(c, app, page)
    items = items_of(int(target.rsplit("/", 1)[1]))
    assert sum(1 for i in items.values() if i.status == "UNVERAENDERT") == 46
    part = file_prices("RaphiLED_Preiserhoehung_Arbeitsscheinwerfer.xlsx", 4, 2)
    for nr, ek in part.items():
        assert items[nr].final_ek == money(str(ek)) and items[nr].final_vk == money(money(str(ek)) * D("2.6"))
    assert items["AS500"].status == "UNVERAENDERT"  # nicht in der Teilliste: bleibt beim zugeordneten Preis


def test_nordlicht_uvp_with_dealer_discount(admin_client, app):
    c = admin_client
    mid = new_manufacturer(c, "Nordlicht Signaltechnik", "NS", "1,8", list_basis="UVP", dealer_discount="35")
    url, page = smart_page(c, "Nordlicht_unsere_Liste_2026.xlsx")
    assert f'<option value="{mid}" selected>' in page
    start(c, app, page)
    url, page = smart_page(c, "Nordlicht_UVP_2027.xlsx")
    assert f'<option value="{mid}" selected>' in page and "= UVP − 35,00 % Händlerrabatt" in page
    target = start(c, app, page)
    items = items_of(int(target.rsplit("/", 1)[1]))
    uvp = file_prices("Nordlicht_UVP_2027.xlsx", 3, 2)
    assert len(items) == 20
    for nr, i in items.items():
        if isinstance(uvp[nr], (int, float)):
            ek = money(D(str(uvp[nr])) * D("0.65"))
            assert (i.final_ek, i.final_vk) == (ek, money(ek * D("1.8"))), nr
            assert i.status == "OK", (nr, i.reasons)
    bad = items["10400"]  # "auf Anfrage": kein Preis -> Fehler, alter Preis bleibt
    assert bad.status == "FEHLER" and bad.decision == "ALT" and "keinen UVP-Preis" in bad.note


def test_svensk_ljus_sek_conversion(admin_client, app):
    c = admin_client
    mid = new_manufacturer(c, "Svensk Ljus AB", "SL", "2,4", list_currency="SEK", exchange_rate="0,095")
    url, page = smart_page(c, "SvenskLjus_unsere_Liste_2026.xlsx")
    start(c, app, page)
    url, page = smart_page(c, "SvenskLjus_Pricelist_2027_SEK.xlsx")
    assert f'<option value="{mid}" selected>' in page
    assert '<option value="SEK" selected>' in page and "aus der Datei" in page and "× Kurs 0,095 in €" in page
    target = start(c, app, page)
    items = items_of(int(target.rsplit("/", 1)[1]))
    sek = file_prices("SvenskLjus_Pricelist_2027_SEK.xlsx", 4, 2)
    assert len(items) == 12
    for nr, i in items.items():
        ek = money(D(str(sek[nr])) * D("0.095"))
        assert (i.final_ek, i.final_vk, i.source_currency, i.status) == (ek, money(ek * D("2.4")), "SEK", "OK"), nr


def test_blocking_problems_are_explained(admin_client, app):
    c = admin_client
    # Herstellerliste ohne eigene Liste und unbekannter Hersteller: nicht startbar, klare Hinweise
    url, page = smart_page(c, "SvenskLjus_Pricelist_2027_SEK.xlsx")
    assert "Hersteller wählen oder neu anlegen" in page and "disabled" in page
    page = c.get(url + "?manufacturer_id=neu").text
    assert "Neuer Hersteller" in page and "noch keine eigene EK/VK-Liste" in page
    # Neuer Hersteller direkt beim Import unserer Liste (mit Faktor)
    url, page = smart_page(c, "SvenskLjus_unsere_Liste_2026.xlsx")
    page = c.get(url + "?manufacturer_id=neu").text
    action, fields = start_form(page)
    r = c.post(action, data={**fields, "new_name": "Svensk Ljus AB", "new_code": "SL"})
    assert r.status_code == 400 and "Faktor" in r.text
    r = c.post(action, data={**fields, "new_name": "Svensk Ljus AB", "new_code": "SL", "new_factor": "2,4"},
               follow_redirects=False)
    assert r.status_code == 303
    run_jobs(app)
    with session_scope() as db:
        m = db.scalar(select(Manufacturer).where(Manufacturer.name == "Svensk Ljus AB"))
        assert m.code == "SL" and m.default_rule_id
        pl = db.scalar(select(PriceList).where(PriceList.kind == "UNSERE"))
        assert pl.manufacturer_id == m.id and pl.status == "IMPORTIERT"
        assert db.scalar(select(Job).order_by(Job.id.desc())).result["kind"] == "UNSERE"
    # SEK-Liste, aber Hersteller auf EUR eingestellt -> blockiert mit Link zu den Einstellungen
    url, page = smart_page(c, "SvenskLjus_Pricelist_2027_SEK.xlsx")
    assert "Die Liste ist in SEK" in page and f'href="/hersteller/{m.id}"' in page and "disabled" in page


def test_lumatec_messy_real_world_list(admin_client, app):
    """Deckblatt, Titelzeilen, Warengruppen-Zwischenzeilen, Preise als Text, EAN/VPE/Staffel-Spalten,
    'entfällt', doppelte Zeile, Fußnoten – so wie Listen wirklich ankommen."""
    c = admin_client
    mid = new_manufacturer(c, "Lumatec Fahrzeugtechnik", "LT", "2,5")
    url, page = smart_page(c, "Lumatec_unsere_Liste_2026.xlsx")
    assert '<option value="UNSERE" selected>' in page and "EK-Preis" not in page.split("Erkannte Spalten")[0][-50:]
    start(c, app, page)

    url, page = smart_page(c, "Lumatec_Preisliste_2027.xlsx")
    assert f'<option value="{mid}" selected>' in page
    assert "Die Spalte „Preis € netto“ wird als EK gelesen" in page  # allgemeine Preisspalte -> laut Hersteller EK
    assert "<b>11 von 11</b>" in page and "1 neue des Herstellers" in page
    assert "disabled" not in page.split("Import &amp; Abgleich starten")[0][-300:]
    target = start(c, app, page)
    items = items_of(int(target.rsplit("/", 1)[1]))
    assert len(items) == 11

    ws = openpyxl.load_workbook(DATA / "Lumatec_Preisliste_2027.xlsx", data_only=True)["Preisliste 2027"]
    new = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[1] and row[5] is not None:
            v = row[5]
            new[row[1]] = money(str(v).replace("€", "").strip().replace(".", "").replace(",", ".")
                                if isinstance(v, str) else str(v))
    for nr, i in items.items():
        if nr == "01003":
            continue
        assert (i.final_ek, i.final_vk, i.status) == (new[nr], money(new[nr] * D("2.5")), "OK"), (nr, i.note)
    gone = items["01003"]
    assert gone.status == "NICHT_IN_HERSTELLERLISTE" and "Laut Hersteller entfallen" in gone.note
    assert "Nachfolger 01004" in gone.note and (gone.final_ek, gone.decision) == (gone.old_amount, "ALT")
    assert items["09001"].status == "OK"  # doppelte Zeile mit gleichem Preis zählt einmal


def test_sales_features_rounding_margin_customer_list(admin_client, app):
    """VK auf ,90 runden, Marge sehen, Abschließen mit Gültig-ab, Preisänderungsliste für Kunden."""
    c = admin_client
    mid = new_manufacturer(c, "Raphi LED", "RA", "2,6")
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "id": str(mid), "name": "Raphi LED", "code": "RA",
                                    "factor": "2,6", "vk_rounding": "e90", "list_basis": "EK",
                                    "review_threshold": "10", "ignore_leading_zeros": "1"})
    assert "Gespeichert" in r.text and 'value="e90" selected' in r.text
    assert "EK × 2,6, auf ,90 (aufrunden)" in c.get("/hersteller").text
    url, page = smart_page(c, "RaphiLED_Preisliste_2026.xlsx")
    start(c, app, page)
    url, page = smart_page(c, "RaphiLED_Herstellerliste_2027.xlsx")
    assert "VK neu = EK × 2,6, auf ,90 (aufrunden)" in page
    target = start(c, app, page)
    items = items_of(int(target.rsplit("/", 1)[1]))
    new = file_prices("RaphiLED_Herstellerliste_2027.xlsx", 4, 2)
    for nr, i in items.items():
        if i.decision == "NEU":
            exact = money(str(new.get(nr, 0))) * D("2.6")
            assert i.final_vk % 1 == D("0.90") and exact <= i.final_vk < exact + 1, (nr, i.final_vk)
            assert i.check_ok is True
    assert "Marge" in c.get(target + "?status=alle").text
    # alles prüfen, abschließen mit Gültig-ab
    for i in items.values():
        if i.status == "NICHT_EINDEUTIG":
            c.post(f"{target}/positionen/{i.id}", data={"csrf_token": c.csrf, "action": "zuordnen", "new_id": "kein"})
    c.post(f"{target}/alle-bestaetigen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    r = c.post(f"{target}/abschliessen", data={"csrf_token": c.csrf, "gueltig_ab": "2027-01-01"})
    assert "gültig ab 01.01.2027" in r.text and "Preisänderungen für Kunden" in r.text
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"{target}/export?art=kunden").content))
    ws = wb["Preisänderungen"]
    assert [h.value for h in ws[1]] == ["Artikelnummer", "Bezeichnung", "Preis bisher", "Preis neu", "Änderung %",
                                        "Gültig ab"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    changed = [i for i in items_of(int(target.rsplit("/", 1)[1])).values() if i.final_vk != i.vk_old]
    assert len(rows) == len(changed) > 0 and all(r[0].startswith("RA") and r[5] == "01.01.2027" for r in rows)
    assert "Marge %" in [h.value for h in openpyxl.load_workbook(io.BytesIO(c.get(f"{target}/export").content))
                         ["Neue Preisliste"][1]]


def _strands_file(path: Path) -> Path:
    """Wie eine echte Strands-Liste: EK in SEK und in Euro nebeneinander, VK in Euro, Nummern mit ST-Kürzel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Artikelnummer", "Bezeichnung", "Serie", "Nettoeinkauf 2025 SEK", "Kurs",
               "Nettoeinkauf 2025 Euro", "Verkauf 2025 Euro"])
    for n in range(1, 21):
        sek = D(100 + n * 7)
        ek = money(sek * D("0.095"))
        ws.append([f"ST27{n:04d}", f"Montagewinkel {n}", "Basic", float(sek), 0.095, float(ek), float(money(ek * 2))])
    wb.save(path)
    return path


def test_our_list_with_ek_in_two_currencies(admin_client, app, tmp_path):
    c = admin_client
    new_manufacturer(c, "Strands", "ST", "2")
    r = upload(c, _strands_file(tmp_path / "Strands EK VK 2025_final.xlsx"))
    url = r.headers["location"]
    page = c.get(url).text
    # EK in Euro wird genommen, die SEK-Spalte bewusst nicht; Währung EUR, startbar
    assert "Unsere Liste braucht eine EK-Spalte" not in page
    assert '<option value="UNSERE" selected>' in page and '<option value="EUR" selected>' in page
    assert "Nettoeinkauf 2025 SEK" in page and "wird nicht verwendet" in page
    assert "disabled" not in page.split("sticky-actions")[1].split("</button>")[0]
    # SEK gewählt: EK-Spalte wechselt auf SEK, VK bleibt Euro -> gemischte Währung blockiert
    page_sek = c.get(url + "?currency=SEK").text
    assert "verschiedene Währungen" in page_sek
    start(c, app, page)
    with session_scope() as db:
        arts = db.scalars(select(Article).join(PriceList).where(PriceList.kind == "UNSERE")).all()
        assert len(arts) == 20
        a = next(x for x in arts if x.article_number.endswith("270001"))
        prices = {pr.price_type: (pr.amount, pr.currency) for pr in a.prices}
        ek = money(D(107) * D("0.095"))
        assert prices == {"EK": (ek, "EUR"), "LISTE": (money(ek * 2), "EUR")}


def test_manufacturer_list_in_two_currencies_follows_setting(admin_client, app, tmp_path):
    c = admin_client
    mid = new_manufacturer(c, "Strands", "ST", "2", list_currency="SEK", exchange_rate="0,095")
    start(c, app, c.get(upload(c, _strands_file(tmp_path / "Strands EK VK 2025.xlsx")).headers["location"]).text)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Art.-Nr.", "Bezeichnung", "Nettoeinkauf 2026 SEK", "Nettoeinkauf 2026 Euro"])
    for n in range(1, 21):
        ws.append([f"27{n:04d}", f"Montagewinkel {n}", 110 + n * 7, 99.0])
    wb.save(tmp_path / "Strands Preisliste 2026.xlsx")
    page = c.get(upload(c, tmp_path / "Strands Preisliste 2026.xlsx").headers["location"]).text
    assert f'<option value="{mid}" selected>' in page and '<option value="HERSTELLER" selected>' in page
    assert '<option value="SEK" selected>' in page and "Nettoeinkauf 2026 Euro" in page
    items = items_of(int(start(c, app, page).rsplit("/", 1)[1]))
    i = items["270001"]
    assert (i.final_ek, i.source_currency) == (money(D(117) * D("0.095")), "SEK")
