"""Jahresabgleich: unsere EK/VK-Liste (Vorjahr) + neue Herstellerliste -> neue EK/VK-Liste."""

import io
import re
from decimal import Decimal

import openpyxl
from sqlalchemy import select

from backend.database.engine import session_scope
from backend.models.entities import AuditLog, Manufacturer, MatchDecision, PriceList, PriceUpdate, PriceUpdateItem
from tests.conftest import make_xlsx, run_jobs
from tests.test_import_web import upload

D = Decimal


def _import(c, app, tmp_path, name, rows, mid, kind, fields=("supplier_price",), currency="EUR"):
    list_id = int(upload(c, make_xlsx(tmp_path / f"{name}.xlsx", rows)).headers["location"].rsplit("/", 1)[1])
    data = {"csrf_token": c.csrf, "sheet": "Preise", "header_row": "1", "col_0": "article_number",
            "col_1": "description", "manufacturer_id": str(mid), "currency": currency, "kind": kind}
    for idx, f in enumerate(fields, start=2):
        data[f"col_{idx}"] = f
        data[f"sep_{idx}"] = ","
    r = c.post(f"/import/{list_id}/confirm", data=data, follow_redirects=False)
    assert r.status_code == 303, r.text
    run_jobs(app)
    return list_id


def _setup(c, factor="2", **mfr):
    c.post("/hersteller", data={"csrf_token": c.csrf, "name": "RaphiLED", "code": "RA", **mfr})
    with session_scope() as db:
        mid = db.scalar(select(Manufacturer)).id
    r = c.post("/regeln/neu", data={"csrf_token": c.csrf, "name": "RaphiLED", "manufacturer_id": str(mid),
                                    "start_price": "EK", "rounding_mode": "HALF_UP", "rounding_places": "2",
                                    "rounding_timing": "STEP", "step_1_type": "multiply", "step_1_value": factor},
               follow_redirects=False)
    rule_id = r.headers["location"].rsplit("/", 1)[1]
    c.post("/hersteller", data={"csrf_token": c.csrf, "id": str(mid), "name": "RaphiLED", "code": "RA",
                                "default_rule_id": rule_id, **mfr})
    return mid, rule_id


def _ours(c, app, tmp_path, mid, rows):
    return _import(c, app, tmp_path, "unsere_2025", [["Artikelnummer", "Bezeichnung", "EK", "VK"], *rows], mid,
                   "UNSERE", ("supplier_price", "list_price"))


def _items(upd_id=None):
    with session_scope() as db:
        upd = db.get(PriceUpdate, upd_id) if upd_id else db.scalars(select(PriceUpdate)).all()[-1]
        items = {i.article_number: i for i in db.scalars(select(PriceUpdateItem)
                                                         .where(PriceUpdateItem.update_id == upd.id))}
        return upd, items


def test_annual_reconciliation_full_flow(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c)
    ours = _ours(c, app, tmp_path, mid, [
        ["LED1", "Modul 1", "10,00", "20,00"],   # +5 % -> OK
        ["LED2", "Modul 2", "10,00", "20,00"],   # +20 % -> prüfen
        ["LED3", "Modul 3", "10,00", "20,00"],   # fehlt beim Hersteller
        ["LED4", "Modul 4", "10,00", "20,00"],   # -10 % genau auf der Schwelle -> prüfen
        ["ABC-1000", "Spezialteil Typ X", "5,00", "10,00"],  # nur unscharfer Treffer
    ])
    new = _import(c, app, tmp_path, "hersteller_2026", [
        ["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "Modul 1", "10,50"], ["LED2", "Modul 2", "12,00"],
        ["LED4", "Modul 4", "9,00"], ["ABC-1001", "Spezialteil Typ X", "6,00"], ["NEU9", "Neuheit", "99,00"]],
        mid, "HERSTELLER")

    # Schritt 5 der Herstellerliste: Abgleich oben, unsere Liste und Regel vorausgewählt
    page = c.get(f"/listen/{new}/kalkulation").text
    assert "Jahresabgleich mit unserer Liste" in page
    assert re.search(rf'<option value="{ours}" selected>', page)
    assert re.search(rf'<option value="{rule_id}" selected>', page)
    assert "Prüfen ab ±10,00 %" in page
    assert "Das ist unsere Liste" in c.get(f"/listen/{ours}/kalkulation").text

    r = c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new,
                                          "rule_id": rule_id}, follow_redirects=False)
    assert r.status_code == 303, r.text
    url = r.headers["location"]
    upd, items = _items()
    assert set(items) == {"LED1", "LED2", "LED3", "LED4", "ABC-1000"}  # NEU9 ignoriert
    assert upd.summary["ignoriert_nur_beim_hersteller"] == 1 and upd.price_type == "EK"

    led1 = items["LED1"]
    assert (led1.status, led1.old_amount, led1.new_amount, led1.vk_old, led1.calculated_amount) == \
        ("OK", D("10.00"), D("10.50"), D("20.00"), D("21.00"))
    assert (led1.final_ek, led1.final_vk, led1.decision, led1.needs_review) == (D("10.50"), D("21.00"), "NEU", False)
    assert led1.difference_percent == D("5.00") and led1.vk_difference_percent == D("5.00")
    led2 = items["LED2"]
    assert led2.status == "PRUEFEN" and led2.reasons == ["EK_AENDERUNG"] and led2.needs_review
    assert led2.final_vk == D("24.00")
    led3 = items["LED3"]
    assert led3.status == "NICHT_IN_HERSTELLERLISTE" and (led3.final_ek, led3.final_vk) == (D("10.00"), D("20.00"))
    assert led3.decision == "ALT" and led3.needs_review
    assert items["LED4"].status == "PRUEFEN" and items["LED4"].difference_percent == D("-10.00")
    abc = items["ABC-1000"]
    assert abc.status == "NICHT_EINDEUTIG" and abc.candidates[0]["number"] == "ABC-1001"
    assert (abc.final_ek, abc.final_vk) == (D("5.00"), D("10.00"))
    assert upd.summary["offen"] == 4 and upd.summary["OK"] == 1

    page = c.get(url).text
    assert 'class="current"><span>6</span>' in page and "RALED1" in page and "Endgültiger Export gesperrt" in page
    page = c.get(url + "?grund=EK_AENDERUNG").text
    assert "RALED2" in page and "RALED4" in page and "RALED1<" not in page and "RALED3" not in page
    page = c.get(url + "?status=offen").text
    assert "RALED3" in page and "RALED1<" not in page

    # Endgültiger Export gesperrt, Entwurf klar gekennzeichnet
    assert c.get(url + "/export").status_code == 409
    wb = openpyxl.load_workbook(io.BytesIO(c.get(url + "/export?entwurf=1").content))
    assert wb.sheetnames == ["Zusammenfassung", "ENTWURF Neue Preisliste", "Zu prüfen"]
    assert "ENTWURF – 4 Positionen ungeprüft" in wb["Zusammenfassung"]["B2"].value

    base = url + "/positionen/"
    # VK von Hand unter EK wird abgelehnt
    r = c.post(base + str(led2.id), data={"csrf_token": c.csrf, "action": "bestaetigen", "manual_vk": "11,00"})
    assert "liegt unter dem EK" in r.text
    r = c.post(base + str(led2.id), data={"csrf_token": c.csrf, "action": "bestaetigen", "manual_vk": "29,90"})
    assert r.status_code == 200
    # Unklare Zuordnung kann nicht einfach bestätigt werden
    r = c.post(base + str(abc.id), data={"csrf_token": c.csrf, "action": "bestaetigen"})
    assert "Zuordnung zuerst auswählen" in r.text
    # Kandidat übernehmen -> neu berechnet (+20 %) und wieder zu prüfen
    c.post(base + str(abc.id), data={"csrf_token": c.csrf, "action": "zuordnen",
                                     "new_id": str(abc.candidates[0]["new_id"])})
    _, items = _items(upd.id)
    assert items["LED2"].decision == "MANUELL" and items["LED2"].final_vk == D("29.90")
    assert items["LED2"].reviewed_by is not None
    abc = items["ABC-1000"]
    assert abc.status == "PRUEFEN" and abc.final_ek == D("6.00") and abc.final_vk == D("12.00")
    assert abc.match_method == "BESTAETIGT" and abc.reviewed_at is None
    with session_scope() as db:
        assert db.scalar(select(MatchDecision)).decision == "MATCH"

    # Sammelbestätigung braucht das Häkchen
    r = c.post(url + "/alle-bestaetigen", data={"csrf_token": c.csrf})
    assert "Häkchen" in r.text
    r = c.post(url + "/alle-bestaetigen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    assert "3 Positionen bestätigt" in r.text
    upd, items = _items(upd.id)
    assert upd.summary["offen"] == 0

    wb = openpyxl.load_workbook(io.BytesIO(c.get(url + "/export").content))
    assert wb.sheetnames == ["Zusammenfassung", "Neue Preisliste", "Zu prüfen"]
    ws = wb["Neue Preisliste"]
    header = [h.value for h in ws[1]]
    rows = {r[0]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {"RALED1", "RALED2", "RALED3", "RALED4", "RAABC-1000"}
    assert (rows["RALED1"]["EK neu"], rows["RALED1"]["VK neu"]) == (10.5, 21)
    assert rows["RALED2"]["VK neu"] == 29.9 and rows["RALED2"]["Entscheidung"] == "VK von Hand"
    assert rows["RALED2"]["geprüft von"] == "admin"
    assert (rows["RALED3"]["EK neu"], rows["RALED3"]["VK neu"]) == (10, 20)
    assert rows["RALED3"]["Entscheidung"] == "alter Preis behalten"
    assert list(wb["Zu prüfen"].iter_rows(min_row=2, values_only=True)) == []
    with session_scope() as db:
        actions = {a.action for a in db.scalars(select(AuditLog))}
    assert {"jahresabgleich", "abgleich_position_bestaetigen", "abgleich_position_zuordnen",
            "abgleich_sammelbestaetigung", "export_entwurf", "export"} <= actions


def test_uvp_list_with_dealer_discount(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c, list_basis="UVP", dealer_discount="40", review_threshold="15")
    ours = _ours(c, app, tmp_path, mid, [["LED1", "Modul 1", "55,00", "110,00"]])
    new = _import(c, app, tmp_path, "rrp", [["Artikelnummer", "Bezeichnung", "UVP"], ["LED1", "Modul 1", "100,00"]],
                  mid, "HERSTELLER", ("rrp",))
    c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new, "rule_id": rule_id})
    upd, items = _items()
    i = items["LED1"]
    assert upd.price_type == "UVP" and upd.dealer_discount == D("40") and upd.review_threshold == D("15")
    # EK = 100 - 40 % = 60 (+9,09 % < 15 %), VK = 60 × 2
    assert (i.source_amount, i.final_ek, i.final_vk, i.status) == (D("100.00"), D("60.00"), D("120.00"), "OK")


def test_vk_below_ek_and_missing_ek_price(admin_client, tmp_path, app):
    c = admin_client
    mid, rule_id = _setup(c, factor="0,9")
    ours = _ours(c, app, tmp_path, mid, [["LED1", "Modul 1", "10,00", "20,00"], ["LED2", "Modul 2", "10,00", "20,00"]])
    new = _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "UVP"], ["LED1", "Modul 1", "10,00"]],
                  mid, "HERSTELLER", ("rrp",))
    c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new, "rule_id": rule_id})
    _, items = _items()
    # Hersteller ist auf EK eingestellt, die Liste hat aber nur UVP -> Fehler, alter Preis bleibt
    i = items["LED1"]
    assert i.status == "FEHLER" and "keinen EK-Preis" in i.note and (i.final_ek, i.final_vk) == (D("10.00"), D("20.00"))
    # FEHLER: kein VK von Hand
    r = c.post(f"/aktualisierungen/{i.update_id}/positionen/{i.id}",
               data={"csrf_token": c.csrf, "action": "bestaetigen", "manual_vk": "30"})
    assert "alte Preis" in r.text

    new2 = _import(c, app, tmp_path, "h2", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "Modul 1", "10,00"]],
                   mid, "HERSTELLER")
    c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new2, "rule_id": rule_id})
    _, items = _items()
    assert items["LED1"].status == "PRUEFEN" and items["LED1"].reasons == ["VK_ABWEICHUNG", "VK_UNTER_EK"]


def test_update_form_validation(admin_client, tmp_path, app):
    c = admin_client
    r = c.post("/aktualisierungen", data={"csrf_token": c.csrf})
    assert r.status_code == 400 and "Herstellerliste wählen" in r.text
    mid, rule_id = _setup(c)
    ours = _ours(c, app, tmp_path, mid, [["LED1", "Modul 1", "10,00", "20,00"]])
    new = _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "Modul 1", "10,00"]],
                  mid, "HERSTELLER")
    # vertauschte Listen
    r = c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": new, "source_id": ours,
                                          "rule_id": rule_id})
    assert r.status_code == 400 and "als Herstellerliste importiert" in r.text
    # ohne Regel kein VK
    r = c.post("/aktualisierungen", data={"csrf_token": c.csrf, "base_id": ours, "source_id": new})
    assert r.status_code == 400 and "Kalkulationsregel" in r.text


def test_manufacturer_settings_validation(admin_client):
    c = admin_client
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "name": "X", "list_basis": "UVP"})
    assert r.status_code == 400 and "Händlerrabatt Pflicht" in r.text
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "name": "X", "list_basis": "EK", "dealer_discount": "120"})
    assert r.status_code == 400 and "zwischen 0 und 100" in r.text
    r = c.post("/hersteller", data={"csrf_token": c.csrf, "name": "X", "review_threshold": "0"})
    assert r.status_code == 400 and "Prüfschwelle" in r.text
    c.post("/hersteller", data={"csrf_token": c.csrf, "name": "X", "list_basis": "UVP", "dealer_discount": "35,5"})
    with session_scope() as db:
        m = db.scalar(select(Manufacturer))
        assert (m.list_basis, m.dealer_discount, m.review_threshold) == ("UVP", D("35.5"), D("10"))


def test_our_list_requires_ek_and_vk(admin_client, tmp_path, app):
    c = admin_client
    mid, _ = _setup(c)
    list_id = int(upload(c, make_xlsx(tmp_path / "x.xlsx", [["Art", "Bez", "EK"], ["A1", "x", "1,00"]]))
                  .headers["location"].rsplit("/", 1)[1])
    r = c.post(f"/import/{list_id}/confirm", data={
        "csrf_token": c.csrf, "sheet": "Preise", "header_row": "1", "col_0": "article_number",
        "col_1": "description", "col_2": "supplier_price", "manufacturer_id": str(mid), "currency": "EUR",
        "kind": "UNSERE"})
    assert r.status_code == 400 and "VK-Spalte" in r.text
    with session_scope() as db:
        assert db.get(PriceList, list_id).kind is None
