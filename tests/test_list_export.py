"""Gespeicherte Listen jederzeit als Excel herunterladen (auch die nach dem Abgleich übernommene Liste)."""

import io
import re

import openpyxl

from tests.test_price_update import _import, _ours, _setup


def _sheet(data: bytes, name="Preisliste"):
    return [list(r) for r in openpyxl.load_workbook(io.BytesIO(data))[name].iter_rows(values_only=True)]


def test_adopted_list_can_be_downloaded_later(admin_client, user_client, tmp_path, app):
    c = admin_client
    mid, _ = _setup(c)
    _ours(c, app, tmp_path, mid, [["LED1", "Leuchte a", "10,00", "20,00"], ["LED2", "Leuchte b", "5,00", "10,00"]])
    _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["LED1", "a", "12,00"]], mid, "HERSTELLER")
    url = c.post(f"/hersteller/{mid}/abgleich", data={"csrf_token": c.csrf}, follow_redirects=False).headers["location"]
    c.post(url + "/alle-bestaetigen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    c.post(url + "/uebernehmen", data={"csrf_token": c.csrf, "bestaetigt": "ja"})
    adopted = int(re.search(r'href="/listen/(\d+)">Neue aktuelle Liste', c.get(url).text).group(1))

    page = c.get(f"/listen/{adopted}").text
    assert f'href="/listen/{adopted}/export"' in page and "Als Excel herunterladen" in page
    assert f'href="/listen/{adopted}/export"' in c.get("/listen").text

    r = c.get(f"/listen/{adopted}/export")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.xlsx"')
    rows = _sheet(r.content)
    assert rows[0][:6] == ["Artikelnummer", "Original-Nr.", "Hersteller", "Bezeichnung", "Kategorie", "Einheit"]
    assert "EK" in rows[0] and "VK" in rows[0] and "Marge %" in rows[0]
    ek, vk = rows[0].index("EK"), rows[0].index("VK")
    by_nr = {row[0]: row for row in rows[1:]}
    assert by_nr["RALED1"][ek] == 12 and by_nr["RALED1"][vk] == 24  # neuer EK × 2
    assert by_nr["RALED2"][ek] == 5 and by_nr["RALED2"][vk] == 10

    # fremde Listen nicht abrufbar
    assert user_client.get(f"/listen/{adopted}/export").status_code == 404


def test_placeholder_ek_has_no_margin_and_kalk_unformatted(admin_client, tmp_path, app):
    from decimal import Decimal

    from backend.excel.export import margin

    assert margin(Decimal("0.01"), Decimal("273.40")) is None
    assert margin(Decimal("84.80"), Decimal("237.44")) == Decimal("64.3")
    c = admin_client
    mid, _ = _setup(c)
    lid = _import(c, app, tmp_path, "kalk", [["Artikelnummer", "Bezeichnung", "EK", "VK", "Kalk"],
                                             ["A1", "ohne Preis", "0,01", "273,40", ""], ["A2", "x", "84,80", "237,44", "2,8"]],
                  mid, "UNSERE", ("supplier_price", "list_price", "calc_factor"))
    rows = _sheet(c.get(f"/listen/{lid}/export").content)
    head = rows[0]
    by_nr = {r[0]: r for r in rows[1:]}
    assert by_nr["RAA1"][head.index("Marge %")] is None
    assert by_nr["RAA2"][head.index("Marge %")] == 64.3
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/listen/{lid}/export").content))
    kalk = [cell for cell in wb["Preisliste"][3] if cell.column == head.index("Kalk") + 1][0]
    assert kalk.value == 2.8 and kalk.number_format == "General"


def test_list_status_flags_placeholder_and_absurd_vk():
    from decimal import Decimal as D

    from backend.excel.export import list_status

    assert list_status("OK", D("0.01"), D("273.40")) == "Preis fehlt (EK 0,01)"
    assert list_status("OK", D("84.80"), D("45766674.00")) == "VK unplausibel (über 20 × EK)"
    assert list_status("OK", D("84.80"), D("259.90")) == "OK"


def test_abgleich_placeholder_and_absurd_old_vk(admin_client, tmp_path, app):
    """EK 0,01 im Vorjahr: zur Prüfung mit „Preis vom Vorjahr fehlt“, alter VK vorgeschlagen, kalkulierter daneben.
    Absurder alter VK (mehr als 20 × EK) wird nie beibehalten."""
    from decimal import Decimal as D

    from tests.test_price_update import _items

    c = admin_client
    mid, _ = _setup(c)  # Standard EK × 2
    _ours(c, app, tmp_path, mid, [["P1", "ohne Vorjahrespreis", "0,01", "200,00"],
                                  ["P2", "absurder VK", "80,00", "45766674,00"],
                                  ["P3", "normal", "80,00", "160,00"]])
    _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["P1", "a", "84,80"], ["P2", "b", "84,80"],
                                    ["P3", "c", "84,80"]], mid, "HERSTELLER")
    c.post(f"/hersteller/{mid}/abgleich", data={"csrf_token": c.csrf}, follow_redirects=False)
    _, items = _items()
    p1, p2, p3 = items["P1"], items["P2"], items["P3"]
    assert p1.status == "PRUEFEN" and "VORJAHR_FEHLT" in p1.reasons and "EK_AENDERUNG" not in p1.reasons
    assert (p1.final_vk, p1.decision, p1.calculated_amount) == (D("200.00"), "MANUELL", D("169.60"))
    assert p1.difference_percent is None and "kalkuliert wären 169,60" in p1.note
    assert p2.status == "PRUEFEN" and "VK_UNPLAUSIBEL" in p2.reasons and "SONDERPREIS" not in p2.reasons
    assert (p2.final_vk, p2.decision) == (D("169.60"), "NEU")
    assert (p3.status, p3.final_vk) == ("OK", D("169.60"))


def test_abgleich_export_factor_compact_and_no_pct_for_placeholder(admin_client, tmp_path, app):
    """Abgleich-Excel: Faktor VK/EK ohne angehängte Nullen (2,8 statt 2,8000), keine EK-Änderung % bei EK 0,01."""
    c = admin_client
    mid, _ = _setup(c, factor="2,8")
    _ours(c, app, tmp_path, mid, [["P1", "ohne Vorjahrespreis", "0,01", "119,90"], ["P2", "normal", "81,60", "228,50"]])
    _import(c, app, tmp_path, "h", [["Artikelnummer", "Bezeichnung", "EK"], ["P1", "a", "47,20"], ["P2", "b", "84,80"]],
            mid, "HERSTELLER")
    url = c.post(f"/hersteller/{mid}/abgleich", data={"csrf_token": c.csrf}, follow_redirects=False).headers["location"]
    wb = openpyxl.load_workbook(io.BytesIO(c.get(url + "/export?entwurf=1").content))
    ws = wb["ENTWURF Neue Preisliste"]
    head = [cell.value for cell in ws[1]]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert rows["RAP1"][head.index("EK Δ%")].value is None
    f = rows["RAP2"][head.index("Faktor VK/EK")]
    assert f.value == 2.8 and f.number_format == "General"
